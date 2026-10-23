"""CSV/XLSX onboarding. Raw files never persist; preview payloads are encrypted."""
import csv
import io
import json
import math
import zipfile
import re
from datetime import datetime, timezone
from uuid import uuid4
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from app.db import db, dumps
from app.core.security import encrypt, decrypt, validate_url, valid_env, cipher
from app.core.prompt_manager import version
from app.schemas import ModelConfig, PromptConfig
from app.services.runs import now
import hashlib
import hmac

_EXCEL_CR_ESCAPE = re.compile(r'_x005F_x005F_|_x005F_x000D_|_x000D_')

def _excel_encode_text(value):
    return (value.replace('_x005F_', '_x005F_x005F_')
                 .replace('_x000D_', '_x005F_x000D_')
                 .replace('\r', '_x000D_'))

def _excel_decode_text(value):
    if not isinstance(value,str):return value
    mapping={'_x005F_x005F_':'_x005F_','_x005F_x000D_':'_x000D_','_x000D_':'\r'}
    return _EXCEL_CR_ESCAPE.sub(lambda found:mapping[found.group()],value)

FIELDS=['team_id','team_name','submission_version','api_base_url','api_key','api_key_env','model_id',
        'is_aggregator','prompt_system','prompt_initial','prompt_review','temperature','max_tokens','provider']
REQUIRED=set(FIELDS[:4]+FIELDS[6:11])


class ImportFormatError(ValueError):
    """Controlled upload error whose message never contains source cell contents."""


def boolean(value):
    if str(value).strip().lower() in ('true','1','yes','是'): return True
    if str(value).strip().lower() in ('false','0','no','否'): return False
    raise ValueError('is_aggregator 必须为 true/false（或是/否）')


def fingerprint(value):
    # Keyed digest prevents offline guessing against persisted submission fingerprints.
    return hmac.new(cipher()._signing_key,dumps(value).encode(),hashlib.sha256).hexdigest()


def read_rows(content, filename):
    if not content or len(content)>8_000_000: raise ValueError('文件不能为空，最大 8MB')
    rows=[]
    if filename.lower().endswith('.csv'):
        try:
            reader=csv.reader(io.StringIO(content.decode('utf-8-sig'),newline=''),strict=True)
            header=next(reader); previous=reader.line_num
            for cells in reader:
                start=previous+1;previous=reader.line_num
                if any(str(v).strip() for v in cells):rows.append((start,cells))
                if len(rows)>100:raise ValueError('每次最多 100 份参赛方案')
        except (UnicodeError,csv.Error,StopIteration):raise ValueError('CSV 解析失败，请使用 UTF-8 标准 CSV；原文不会回显') from None
    elif filename.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist())>40_000_000:raise ImportFormatError('Excel 解压大小超过限制')
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
            sheet=book['submissions'] if 'submissions' in book.sheetnames else book.worksheets[0]
            encoded_lines=book.properties.keywords == 'wargame-crlf-preserved-v1'
            iterator=sheet.iter_rows(); header=[cell.value for cell in next(iterator)]
            for number,cells in enumerate(iterator,2):
                if any(c.value is not None for c in cells):
                    if any(c.data_type=='f' for c in cells):raise ImportFormatError(f'第 {number} 行包含 Excel 公式，请改为纯文本')
                    rows.append((number,[_excel_decode_text(c.value) if encoded_lines else c.value for c in cells]))
                if len(rows)>100:raise ImportFormatError('每次最多 100 份参赛方案')
            book.close()
        except ImportFormatError:raise
        except Exception:raise ValueError('Excel 解析失败，仅支持有效 .xlsx；原文不会回显') from None
    else:raise ValueError('只支持 .xlsx 或 UTF-8 .csv')
    header=[str(h or '').strip() for h in header]
    if len(set(header))!=len(header) or not REQUIRED<=set(header) or set(header)-set(FIELDS):
        raise ValueError('表头缺失、重复或包含未知字段，请使用系统模板；api_key 与 api_key_env 至少提供一列')
    if not {'api_key','api_key_env'}&set(header):raise ValueError('缺少凭据列')
    if not rows:raise ValueError('没有参赛数据')
    return header,rows


def preview_import(content,filename):
    header,rows=read_rows(content,filename)
    with db() as c:
        existing={(r['team_id'],r['version']):r['fingerprint'] for r in c.execute('SELECT team_id,version,fingerprint FROM submissions')}
        team_names={r['id']:r['name'] for r in c.execute('SELECT * FROM teams')}
    seen={};valid=[];items=[]
    for number,cells in rows:
        item={'row':number,'status':'error','errors':[]}
        if len(cells)!=len(header):
            item['errors']=['列数与表头不一致'];items.append(item);continue
        raw={h:('' if v is None else str(v)) for h,v in zip(header,cells)}
        def text(key):return raw.get(key,'').strip()
        try:
            for key in ('team_id','team_name','submission_version','model_id'):
                if not text(key) or len(text(key))>80:raise ValueError('队伍编号、名称、版本、模型名称须为 1～80 个字符')
            prompts={key:raw.get('prompt_'+key,'') for key in ('system','initial','review')}
            if any(not p.strip() or len(p)>20000 for p in prompts.values()):raise ValueError('三段 Prompt 须非空且每段不超过 20000 字符')
            provider=text('provider') or 'openai'
            key,env=text('api_key'),text('api_key_env')
            if key and env:raise ValueError('api_key 与 api_key_env 只能填写一个')
            if not valid_env(env):raise ValueError('凭据环境变量名不合法')
            if provider=='openai' and not(key or env):raise ValueError('真实模型必须提供 API Key 或环境变量引用')
            if provider=='mock' and (key or env or text('api_base_url')):raise ValueError('Mock 请勿填写真实 URL 或凭据')
            url=validate_url(text('api_base_url')) if provider=='openai' else ''
            try:
                model=ModelConfig(name=text('team_name'),provider=provider,base_url=url,model_id=text('model_id'),
                                  api_key=key,api_key_env=env,is_aggregator=boolean(text('is_aggregator')),
                                  temperature=float(text('temperature') or 0),max_tokens=int(text('max_tokens') or 1024)).model_dump()
            except Exception:raise ValueError('模型类型、聚合声明或推理参数无效；温度 0～2，max_tokens 为 1～16384 的整数') from None
            tid,ver=text('team_id'),text('submission_version')
            normalized={'team_id':tid,'name':text('team_name'),'version':ver,'model':model,'prompts':prompts,'row':number}
            fp=fingerprint({k:v for k,v in normalized.items() if k!='row'})
            if (tid,ver) in seen:
                item['status']='duplicate' if seen[(tid,ver)]==fp else 'error'
                if item['status']=='error':item['errors']=['同一文件相同队伍与版本包含不同内容']
            elif (tid,ver) in existing:
                item['status']='duplicate' if existing[(tid,ver)]==fp else 'error'
                if item['status']=='error':item['errors']=['该队伍版本已存在且内容不同，请使用新版本号']
            else:
                if tid in team_names and team_names[tid]!=normalized['name']:raise ValueError('相同 team_id 的队伍名称不一致')
                team_names[tid]=normalized['name'];item['status']='ready';normalized['fingerprint']=fp;valid.append(normalized)
            seen[(tid,ver)]=fp
            item.update(team_id=tid,team_name=normalized['name'],submission_version=ver,model_id=model['model_id'],
                        api_base_url=url,provider=provider,is_aggregator=model['is_aggregator'],prompts_complete=True,
                        credential_provided=bool(key or env))
        except ValueError as exc:
            # Only our controlled messages and validate_url's generic error, never raw parser input.
            item['errors']=[str(exc)]
        items.append(item)
    identity=uuid4().hex
    summary={'preview_id':identity,'total':len(items),'ready':sum(i['status']=='ready' for i in items),
             'errors':sum(i['status']=='error' for i in items),'duplicates':sum(i['status']=='duplicate' for i in items),
             'items':items,'api_requests':0}
    payload={'rows':valid,'summary':summary}
    with db() as c:
        c.execute("UPDATE import_previews SET secret=NULL WHERE result IS NULL AND datetime(created_at) < datetime('now','-1 hour')")
        c.execute('INSERT INTO import_previews VALUES(?,?,?,NULL)',(identity,encrypt(dumps(payload)),now()))
    return summary


def confirm_import(identity):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        row=c.execute('SELECT * FROM import_previews WHERE id=?',(identity,)).fetchone()
        if not row:raise ValueError('导入预览不存在')
        if row['result']:return json.loads(row['result'])
        if (datetime.now(timezone.utc)-datetime.fromisoformat(row['created_at'])).total_seconds()>3600 or not row['secret']:
            raise ValueError('导入预览已过期，请重新上传')
        payload=json.loads(decrypt(row['secret']))
        if payload['summary']['errors']:raise ValueError('预览含错误，整批不导入；请修正后重新上传')
        created=[];duplicates=payload['summary']['duplicates'];selection_updates=[]
        saved=c.execute("SELECT body FROM organizer_state WHERE key='official_roster'").fetchone()
        if saved:
            roster=set(json.loads(saved['body']))
        else:
            roster=set();seen=set()
            for old in c.execute("SELECT s.id,s.team_id,m.body AS model FROM submissions s JOIN models m ON m.id=s.model_id ORDER BY s.team_id,s.created_at DESC,s.rowid DESC"):
                if old['team_id'] not in seen and json.loads(old['model'])['provider']!='mock':
                    roster.add(old['id']);seen.add(old['team_id'])
        for r in payload['rows']:
            old=c.execute('SELECT id,fingerprint FROM submissions WHERE team_id=? AND version=?',(r['team_id'],r['version'])).fetchone()
            if old:
                if old['fingerprint']!=r['fingerprint']:raise ValueError('并发导入产生版本冲突，本次事务已回滚')
                duplicates+=1;continue
            team=c.execute('SELECT name FROM teams WHERE id=?',(r['team_id'],)).fetchone()
            if team and team['name']!=r['name']:raise ValueError('队伍名称发生冲突，本次事务已回滚')
            mid,pid,sid=uuid4().hex,uuid4().hex,uuid4().hex
            model=dict(r['model']);key=model.pop('api_key')
            prompt=PromptConfig(name=r['name'],**r['prompts'],default_calls=1).model_dump();prompt['version']=version(prompt)
            c.execute('INSERT OR IGNORE INTO teams VALUES(?,?)',(r['team_id'],r['name']))
            c.execute('INSERT INTO models VALUES(?,?,?)',(mid,dumps(model),encrypt(key)))
            c.execute('INSERT INTO prompts VALUES(?,?)',(pid,dumps(prompt)))
            c.execute('INSERT INTO submissions VALUES(?,?,?,?,?,?,?,?)',(sid,r['team_id'],r['version'],r['name'],mid,pid,r['fingerprint'],now()))
            created.append(sid)
            if model['provider']!='mock':
                previous=None
                for old in c.execute('SELECT id,version FROM submissions WHERE team_id=?',(r['team_id'],)).fetchall():
                    if old['id'] in roster and old['id']!=sid:previous=old['version']
                    roster.discard(old['id'])
                roster.add(sid)
                # previous_version is None when the team was not on the roster before this import.
                selection_updates.append({'team_id':r['team_id'],'team_name':r['name'],
                                          'submission_version':r['version'],'submission_id':sid,
                                          'previous_version':previous})
        if selection_updates:
            c.execute("INSERT INTO organizer_state(key,body) VALUES('official_roster',?) ON CONFLICT(key) DO UPDATE SET body=excluded.body",(dumps(sorted(roster)),))
        result={'imported':len(created),'duplicates':duplicates,'submission_ids':created,
                'api_requests':0,'selection_updates':selection_updates}
        c.execute('UPDATE import_previews SET secret=NULL,result=? WHERE id=?',(dumps(result),identity))
    return result


def roster_ids(c):
    """Return (configured, submission_ids): saved organizer roster, else newest real version per team."""
    row=c.execute("SELECT body FROM organizer_state WHERE key='official_roster'").fetchone()
    if row:return True,json.loads(row['body'])
    ids=[];seen=set()
    for r in c.execute("""SELECT s.id,s.team_id,m.body AS model FROM submissions s
        JOIN models m ON m.id=s.model_id ORDER BY s.team_id,s.created_at DESC,s.rowid DESC"""):
        if r['team_id'] not in seen and json.loads(r['model'])['provider']!='mock':
            ids.append(r['id']);seen.add(r['team_id'])
    return False,ids


def list_submissions():
    with db() as c:
        rows=c.execute('SELECT s.*,m.body AS model,p.body AS prompt,m.secret FROM submissions s LEFT JOIN models m ON m.id=s.model_id '
                       'LEFT JOIN prompts p ON p.id=s.prompt_id ORDER BY s.team_id,s.created_at DESC').fetchall()
    output=[];seen=set();latest={}
    for r in rows:
        model=json.loads(r['model']) if r['model'] else {};prompt=json.loads(r['prompt']) if r['prompt'] else {}
        is_latest=r['team_id'] not in seen
        seen.add(r['team_id'])
        if is_latest:latest[r['team_id']]=r['version']
        output.append({'id':r['id'],'team_id':r['team_id'],'team_name':r['name'],'submission_version':r['version'],
                       'model_id':model.get('model_id'),'provider':model.get('provider'),'api_base_url':model.get('base_url'),
                       'is_aggregator':model.get('is_aggregator'),'credential_provided':bool(r['secret'] or model.get('api_key_env')),
                       'prompt_version':prompt.get('version'),'available':bool(model and prompt),
                       'is_latest':is_latest,'latest_version':latest.get(r['team_id']),
                       'credential_source':'环境变量' if model.get('api_key_env') else ('API Key' if r['secret'] else '未提供')})
    return output


def template_rows(mock=False):
    return [[f'team-{i}',f'Mock 队伍 {i}' if mock else f'参赛队伍 {i}','v1','' if mock else 'https://api.example.com/v1','',
             '' if mock else f'TEAM_{i}_API_KEY','mock-demo' if mock else 'MODEL_ID','false',
             '仅根据当前截图和背景判断。\n遵循原题输出要求。','请核对阵营、位置、旗标。\n不要补充未给出的信息，保持独立作答。',
             '请复核上一轮回答。\n检查遗漏与矛盾，有误则修正。',0,1024,'mock' if mock else 'openai'] for i in range(1,6)]


def spreadsheet_bytes(headers,rows,sheet_name='submissions'):
    book=Workbook();book.properties.keywords='wargame-crlf-preserved-v1' if sheet_name=='submissions' else '';sheet=book.active;sheet.title=sheet_name;sheet.append(headers)
    for row in rows:
        sheet.append([_excel_encode_text(v) if sheet_name=='submissions' and isinstance(v,str) else v for v in row])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
        sheet.row_dimensions[sheet.max_row].height=85
    for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='193957');cell.font=Font(color='FFFFFF',bold=True)
    from openpyxl.utils import get_column_letter
    for i,h in enumerate(headers,1):sheet.column_dimensions[get_column_letter(i)].width=48 if 'prompt' in h else 24
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    stream=io.BytesIO();book.save(stream);return stream.getvalue()


def csv_bytes(headers,rows):
    stream=io.StringIO(newline='');writer=csv.writer(stream);writer.writerow(headers)
    for row in rows:writer.writerow(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')) else v for v in row])
    return ('\ufeff'+stream.getvalue()).encode('utf-8')
