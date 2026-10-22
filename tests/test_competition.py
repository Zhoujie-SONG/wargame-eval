import asyncio
import copy
import io
import json
import socket
import time
from concurrent.futures import ThreadPoolExecutor
import httpx
import pytest
from openpyxl import load_workbook
from app.db import db,dumps,init_db
from app.services import submissions as subs


def upload(client,format='csv',rows=None):
    rows=subs.template_rows(True) if rows is None else rows
    content=subs.csv_bytes(subs.FIELDS,rows) if format=='csv' else subs.spreadsheet_bytes(subs.FIELDS,rows)
    r=client.post('/api/submissions/batch/preview',files={'file':('teams.'+format,content)})
    assert r.status_code==200,r.text
    return r.json()


def import_five(client,format='csv',rows=None):
    p=upload(client,format,rows)
    assert p['errors']==0,p
    r=client.post('/api/submissions/batch/import',json={'preview_id':p['preview_id']})
    assert r.status_code==200,r.text
    return r.json()['submission_ids']


def batch_payload(ids,**kwargs):
    p={'request_key':'test-key-'+str(time.monotonic_ns()),'name':'五队联调','submission_ids':ids,'calls':1,
       'question_ids':['G01-01-CORE','G01-02-CORE','G01-01-P1','G01-01-C1']}
    p.update(kwargs)
    return p


@pytest.mark.parametrize('scope',[{'limit':1},{'limit':1279},{'dimension':'L2_main'},
                                  {'question_ids':['G01-01-CORE']}])
def test_official_rejects_explicit_scope_before_preview_or_save(imported,scope):
    c=imported;ids=import_five(c)
    p=batch_payload(ids,official=True,question_ids=[],**{k:v for k,v in scope.items() if k!='question_ids'})
    p.update(scope)
    for endpoint in ('/api/evaluation-batches/preview','/api/evaluation-batches'):
        r=c.post(endpoint,json=p)
        assert r.status_code==400 and '不允许' in r.json()['detail']
    assert c.get('/api/evaluation-batches').json()==[]
    assert c.get('/api/runs').json()==[]


def test_official_mock_missing_images_full_gate_and_confirmation(imported,monkeypatch):
    c=imported;ids=import_five(c)
    p=batch_payload(ids[:1],official=True,question_ids=[])
    for endpoint in ('/api/evaluation-batches/preview','/api/evaluation-batches'):
        r=c.post(endpoint,json=p)
        assert r.status_code==400 and 'Mock' in r.text
    monkeypatch.setattr(socket,'getaddrinfo',lambda *a,**k:[(2,1,6,'',('93.184.216.34',443))])
    rows=subs.template_rows(False)[:1];rows[0][0]='real-team';rows[0][4]='fixture-only-secret';rows[0][5]=''
    real=import_five(c,rows=rows)
    p=batch_payload(real,official=True,question_ids=[])
    for endpoint in ('/api/evaluation-batches/preview','/api/evaluation-batches'):
        r=c.post(endpoint,json=p)
        assert r.status_code==400 and '真实截图' in r.text
    from test_system import png
    from app.db import question_rows
    image=c.post('/api/images',files={'file':('fixture.png',png('gold'))},data={'real':'true'}).json()
    questions=question_rows()
    mappings=[dict(replay_id=r,image_step=s,image_id=image['id']) for r,s in {(q['replay_id'],q['image_step']) for q in questions}]
    assert c.post('/api/images/map',json=mappings).status_code==200
    # Even an explicit list containing the whole dataset is prohibited in official mode.
    r=c.post('/api/evaluation-batches/preview',json=dict(p,question_ids=[q['question_id'] for q in questions]))
    assert r.status_code==400
    preview=c.post('/api/evaluation-batches/preview',json=p)
    assert preview.status_code==200,preview.text
    assert preview.json()['official'] and preview.json()['questions_per_team']==1279
    assert preview.json()['mode_label']=='正式比赛｜全量 1,279 道 QA'
    bid=c.post('/api/evaluation-batches',json=p).json()['batch_id']
    assert c.get('/api/evaluation-batches/'+bid).json()['status']=='draft'
    queued=[]
    from app.services import competition
    monkeypatch.setattr(competition.runner,'enqueue_many',queued.extend)
    assert c.post('/api/evaluation-batches/'+bid+'/start').status_code==400
    assert c.post('/api/evaluation-batches/'+bid+'/start',json={'confirm_official':False}).status_code==400
    assert not queued
    assert c.post('/api/evaluation-batches/'+bid+'/start',json={'confirm_official':True}).status_code==200
    assert len(queued)==1  # Queue stub: no real API requests are made.
    assert c.get('/api/evaluation-batches/'+bid+'/analytics').json()['official']


def test_test_mode_partial_classification(imported):
    c=imported;ids=import_five(c)
    b=run_batch(c,ids[:1],official=False,dimension='L2_main',limit=1,question_ids=[])
    assert b['completed']==1 and not b['official']
    assert b['mode_label']=='历史测试｜非正式成绩'
    assert b['external_api_calls']==0 and b['runs'][0]['kind']=='mock'
    assert c.get('/api/leaderboard').json()==[]
    d=c.get('/api/evaluation-batches/'+b['batch_id']+'/analytics').json()
    assert not d['official'] and len(d['groups'])==1


def seed_legacy_batch(c,payload):
    """Construct an old-format test record in the isolated DB; never bypass production admission."""
    from uuid import uuid4
    from app.schemas import RunConfig
    from app.services.runs import prepare_run,now
    bid=uuid4().hex;ids=[];plans=[]
    with db() as con:
        submissions=[dict(con.execute('SELECT * FROM submissions WHERE id=?',(sid,)).fetchone()) for sid in payload['submission_ids']]
    for sub in submissions:
        fields={k:payload[k] for k in ('mode','calls','early_stop','dimension','limit','question_ids') if k in payload}
        body,secret=prepare_run(RunConfig(name=sub['name'],model_id=sub['model_id'],prompt_id=sub['prompt_id'],official=False,**fields))
        rid=uuid4().hex;ids.append(rid);body.update(batch_id=bid,team_id=sub['team_id'],submission_version=sub['version']);plans.append((sub,rid,body,secret))
    count=len(plans[0][2]['questions']);calls=1 if payload.get('mode','single_pass')=='single_pass' else payload['calls']
    summary=dict(official=False,teams=len(plans),questions_per_team=count,total_qa=len(plans)*count,max_requests=len(plans)*count*calls,max_external_requests=0,mode=payload.get('mode','single_pass'),calls=calls)
    with db() as con:
        con.execute('INSERT INTO evaluation_batches(id,request_key,name,created_at,body) VALUES(?,?,?,?,?)',(bid,payload['request_key'],payload['name'],now(),dumps({'request':payload,'plan':summary})))
        for sub,rid,body,secret in plans:
            con.execute('INSERT INTO runs(id,status,created_at,body,secret) VALUES(?,?,?,?,?)',(rid,'draft',now(),dumps(body),secret))
            con.execute('INSERT INTO batch_entries VALUES(?,?,?,?)',(bid,sub['team_id'],sub['id'],rid))
    return bid


def execute_legacy_fixture(c,bid):
    """Exercise the unchanged worker with historical fixture runs; public start must reject them."""
    from app.services.competition import runner
    with db() as con:
        ids=[r[0] for r in con.execute('SELECT run_id FROM batch_entries WHERE batch_id=?',(bid,))]
        con.execute('UPDATE evaluation_batches SET started=1 WHERE id=?',(bid,))
        for rid in ids:con.execute("UPDATE runs SET status='queued' WHERE id=?",(rid,))
    c.portal.call(runner.enqueue_many,ids)


def wait_batch(c,bid):
    for _ in range(1000):
        b=c.get('/api/evaluation-batches/'+bid).json()
        if b['status'] not in ('draft','running'):return b
        time.sleep(.01)
    raise AssertionError('batch not completed')


def run_batch(c,ids,**kwargs):
    bid=seed_legacy_batch(c,batch_payload(ids,**kwargs));execute_legacy_fixture(c,bid)
    return wait_batch(c,bid)


@pytest.mark.parametrize('format',['csv','xlsx'])
def test_five_teams_multiline_and_idempotent(client,format):
    rows=subs.template_rows(True)
    rows[0][8]='中文逗号, "引号"\n第二行\r\n第三行'
    p=upload(client,format,rows)
    assert p['total']==5 and p['ready']==5 and p['api_requests']==0
    assert client.get('/api/submissions').json()==[]
    result=client.post('/api/submissions/batch/import',json={'preview_id':p['preview_id']}).json()
    assert result['imported']==5
    assert client.post('/api/submissions/batch/import',json={'preview_id':p['preview_id']}).json()==result
    assert upload(client,format,rows)['duplicates']==5
    with db() as con:
        saved=json.loads(con.execute('SELECT p.body FROM prompts p JOIN submissions s ON s.prompt_id=p.id WHERE s.team_id=?',('team-1',)).fetchone()[0])
        assert saved['system']==rows[0][8]
        assert con.execute('SELECT count(*) FROM submissions').fetchone()[0]==5
        assert all(r[0] is None for r in con.execute('SELECT secret FROM import_previews WHERE result IS NOT NULL'))


def test_preview_errors_and_transaction(client,monkeypatch):
    rows=subs.template_rows(True);rows[1][8]='';rows[2][13]='wrong';rows[3][3]='https://127.0.0.1';rows[3][13]='openai';rows[3][4]='DO-NOT-LEAK'
    p=upload(client,rows=rows)
    assert p['errors']==3 and p['items'][1]['row']==6  # CSV physical line, including previous multiline fields
    assert 'DO-NOT-LEAK' not in json.dumps(p)
    assert client.post('/api/submissions/batch/import',json={'preview_id':p['preview_id']}).status_code==400
    assert client.get('/api/submissions').json()==[]
    rows=subs.template_rows(True);rows[0][9]='different';first=import_five(client,rows=rows)
    rows[0][8]='modified same version'
    assert upload(client,rows=rows)['errors']==1
    rows[0][2]='v2'
    new=client.post('/api/submissions/batch/import',json={'preview_id':upload(client,rows=rows)['preview_id']}).json()
    assert new['imported']==1 and len(client.get('/api/submissions').json())==6


def test_credentials_encrypted_and_not_echoed(client,monkeypatch):
    monkeypatch.setattr(socket,'getaddrinfo',lambda *a,**k:[(2,1,6,'',('93.184.216.34',443))])
    rows=subs.template_rows(False);rows[0][4]='sk-unique-secret-for-test';rows[0][5]=''
    p=upload(client,rows=rows)
    assert p['errors']==0
    client.post('/api/submissions/batch/import',json={'preview_id':p['preview_id']})
    for path in ('/api/submissions','/api/models','/api/evaluation-batches'):
        assert rows[0][4] not in client.get(path).text
    with db() as c:
        assert rows[0][4] not in ' '.join(r[0] or '' for r in c.execute('SELECT secret FROM models'))
        assert rows[0][4] not in ' '.join(r[0] or '' for r in c.execute('SELECT secret FROM import_previews'))
    assert client.get('/api/runs').json()==[]


def test_template_files(client):
    for fmt in ('csv','xlsx'):
        r=client.get('/api/submissions/template?mock=true&format='+fmt)
        assert r.status_code==200
        p=client.post('/api/submissions/batch/preview',files={'file':('mock.'+fmt,r.content)})
        assert p.status_code==200 and p.json()['ready']==5
    xlsx=client.get('/api/submissions/template?format=xlsx').content
    b=load_workbook(io.BytesIO(xlsx));assert b.active.max_row==6 and b.active.freeze_panes=='A2'


def test_parser_exception_does_not_echo_cells(client,monkeypatch):
    content=subs.spreadsheet_bytes(subs.FIELDS,subs.template_rows(True))
    def broken(*args,**kwargs):raise ValueError('malformed numeric cell: sk-private-source-secret')
    monkeypatch.setattr(subs,'load_workbook',broken)
    r=client.post('/api/submissions/batch/preview',files={'file':('bad.xlsx',content)})
    assert r.status_code==400 and 'sk-private-source-secret' not in r.text


@pytest.mark.parametrize('mode,calls',[('single_pass',1),('self_refine',3),('majority_vote',3),('self_refine',5),('majority_vote',5)])
def test_full_batch_pipeline(imported,mode,calls):
    c=imported;ids=import_five(c)
    if calls>3:assert c.put('/api/settings',json={'max_calls_per_question':calls,'timeout_per_question_seconds':60}).status_code==200
    payload=batch_payload(ids,mode=mode,calls=calls)
    assert c.post('/api/evaluation-batches/preview',json=dict(payload,official=False)).status_code==400
    assert c.post('/api/evaluation-batches',json=dict(payload,official=False)).status_code==400
    bid=seed_legacy_batch(c,payload)
    before=c.get('/api/evaluation-batches/'+bid).json()
    assert before['status']=='draft' and before['actual_calls']==0
    assert c.post('/api/evaluation-batches/'+bid+'/start').status_code==400
    execute_legacy_fixture(c,bid)
    b=wait_batch(c,bid)
    assert b['status']=='completed' and b['finished']==5 and b['completed']==20
    assert b['actual_calls']==20*calls and b['external_api_calls']==0
    data=c.get('/api/evaluation-batches/'+bid+'/analytics').json()
    assert len(data['groups'])==1 and len(data['groups'][0]['entries'])==5
    assert all(r['kind']=='mock' for r in data['groups'][0]['entries'])
    assert c.get('/api/leaderboard').json()==[]
    assert sum(r['completed'] for r in b['runs'])==b['completed']
    assert all(r['mean_question_seconds'] is not None for r in data['groups'][0]['entries'])
    for fmt in ('html','json','csv','xlsx'):
        for audience in ('public','admin'):
            r=c.get(f'/api/evaluation-batches/{bid}/export?format={fmt}&audience={audience}')
            assert r.status_code==200,r.text
            if fmt!='xlsx':
                assert 'api_key' not in r.text and 'api_key_env' not in r.text and 'gt_answer' not in r.text
                assert '仅根据当前截图' not in r.text
            if fmt=='html':
                assert r.text.count('<svg ')==4 and 'src="https://' not in r.text and 'href="https://' not in r.text
                assert 'Mock' in r.text and '平均单题' in r.text
                assert "style-src 'unsafe-inline'" in r.headers['content-security-policy']
    group=data['groups'][0]['id']
    for chart in ('ranking','dimensions','radar','efficiency'):
        svg=c.get(f'/api/evaluation-batches/{bid}/charts/{chart}.svg?group={group}').text
        import xml.etree.ElementTree as ET
        assert ET.fromstring(svg).tag.endswith('svg')


def test_missing_images_official_and_budget_gate(imported):
    c=imported;ids=import_five(c)
    assert c.post('/api/evaluation-batches',json=batch_payload(ids,official=True)).status_code==400
    assert c.post('/api/evaluation-batches',json=batch_payload(ids,mode='majority_vote',calls=4)).status_code==400
    assert c.post('/api/evaluation-batches',json=batch_payload(ids,calls=3)).status_code==400
    assert c.post('/api/evaluation-batches',json=batch_payload(ids,submission_ids=ids+[ids[0]])).status_code==400
    assert c.get('/api/evaluation-batches').json()==[]


def test_failed_model_does_not_stop_others(imported,monkeypatch):
    c=imported;ids=import_five(c)
    import app.services.runs as runs
    original=runs.make_client
    class Broken:
        async def complete(self,*args):raise httpx.ConnectError('SECRET_PROVIDER_MESSAGE')
    monkeypatch.setattr(runs,'make_client',lambda m,k:Broken() if m['name']=='Mock 队伍 2' else original(m,k))
    b=run_batch(c,ids)
    assert b['finished']==5 and b['error_count']==4
    assert sum(r['error_count']==0 for r in b['runs'])==4 and b['actual_calls']==20
    assert 'SECRET_PROVIDER_MESSAGE' not in c.get('/api/evaluation-batches/'+b['batch_id']+'/export?format=json&audience=admin').text


def test_stop_and_persistent_recovery(imported,monkeypatch):
    c=imported;ids=import_five(c)
    import app.services.runs as runs
    class Slow:
        async def complete(self,*args):
            await asyncio.sleep(.1);return {'output':'A','usage':{}}
    monkeypatch.setattr(runs,'make_client',lambda *a:Slow())
    bid=seed_legacy_batch(c,batch_payload(ids,limit=20,question_ids=[]))
    execute_legacy_fixture(c,bid);time.sleep(.02)
    assert c.post('/api/evaluation-batches/'+bid+'/stop').status_code==200
    b=wait_batch(c,bid)
    assert b['actual_calls']<=1 and b['status']=='stopped_or_interrupted'
    assert c.get('/api/evaluation-batches/'+bid+'/leaderboard').json()['groups']==[]
    other=seed_legacy_batch(c,batch_payload(ids))
    with db() as con:
        rid=con.execute('SELECT run_id FROM batch_entries WHERE batch_id=?',(other,)).fetchone()[0]
        con.execute("UPDATE runs SET status='running' WHERE id=?",(rid,))
    init_db()
    assert c.get('/api/runs/'+rid).json()['status']=='interrupted'
    assert c.get('/api/evaluation-batches/'+other).json()['failed_or_interrupted']==1


def test_filter_fairness_rerun_and_xss(imported):
    c=imported;rows=subs.template_rows(True);rows[0][1]='<script>alert("x")</script>'
    ids=import_five(c,rows=rows);b=run_batch(c,ids);bid=b['batch_id']
    d=c.get('/api/evaluation-batches/'+bid+'/analytics?team_id=team-1&top=5').json()
    assert len(d['groups'][0]['entries'])==1 and d['selected_teams']==1
    h=c.get('/api/evaluation-batches/'+bid+'/export?team_id=team-1').text
    assert '<script>alert' not in h and '&lt;script&gt;' in h
    assert c.get('/api/evaluation-batches/'+bid+'/analytics?team_id=unknown').status_code==400
    # Simulate an historical differing-mode record: it must become a separate group.
    with db() as con:
        rid=b['runs'][0]['run_id'];r=con.execute('SELECT body FROM runs WHERE id=?',(rid,)).fetchone()
        body=json.loads(r[0]);body['mode']='majority_vote'
        con.execute('UPDATE runs SET body=? WHERE id=?',(dumps(body),rid))
    assert len(c.get('/api/evaluation-batches/'+bid+'/analytics').json()['groups'])==2
    qid=c.get('/api/runs/'+rid).json()['results'][0]['question_id']
    c.post('/api/runs/'+rid+'/infrastructure',json={'question_ids':[qid],'reason':'主办方资源测试故障'})
    replacement=c.post('/api/runs/'+rid+'/rerun').json()['run_id']
    for _ in range(500):
        if c.get('/api/runs/'+replacement).json()['status']=='completed':break
        time.sleep(.01)
    runs=c.get('/api/evaluation-batches/'+bid).json()['runs']
    assert len(runs)==5 and replacement in [r['run_id'] for r in runs] and rid not in [r['run_id'] for r in runs]


def test_blank_chart_and_timing_not_zero(imported):
    c=imported;ids=import_five(c)
    bid=seed_legacy_batch(c,batch_payload(ids))
    html=c.get('/api/evaluation-batches/'+bid+'/export').text
    assert '暂无可排名成绩' in html and '<svg' not in html
    assert len(c.get('/api/evaluation-batches/'+bid+'/leaderboard').json()['excluded'])==5
    inline=c.get('/api/evaluation-batches/'+bid+'/export?format=html&download=false')
    assert inline.headers['content-disposition'].startswith('inline')
